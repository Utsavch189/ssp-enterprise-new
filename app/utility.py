import bcrypt 
from dotenv import load_dotenv
import os
import uuid
import time
from openpyxl import load_workbook
import functools
from flask import jsonify,make_response,redirect,request,render_template
import json
from datetime import datetime
import sib_api_v3_sdk
from sib_api_v3_sdk.rest import ApiException
from sib_api_v3_sdk.models import SendSmtpEmail
import requests
from app.db import get_con_cursor
from multiprocessing import Process
from premailer import transform

import smtplib
import base64
from email import encoders
from email.mime.text import MIMEText
from email.mime.multipart import MIMEMultipart
from email.mime.base import MIMEBase

load_dotenv()

salt=os.getenv("SALT")

def is_authorize(redirect_url="/login",json_response=False,role=None):
    def inner(func):
        @functools.wraps(func)
        def wrapper(*args, **kwargs):
            session_data = request.session_data

            if json_response:
                response = make_response(jsonify({"message":"unauthorized!","status":401, "redirect":redirect_url}),401)
            else:
                response = redirect(redirect_url)

            if not session_data:
                return response

            session_data = json.loads(session_data)
            expires_at = datetime.strptime(session_data.get('expires_at'), "%Y-%m-%d %H:%M:%S.%f")

            if expires_at < datetime.now():
                return response
            
            if role and session_data.get('role')!=role:
                return response

            result = func(*args, **kwargs,session_data=session_data)
            return result

        return wrapper
    return inner

def download_save_image(image_name: str, image_url: str, folder: str = "app/media"):
    try:
        # Check if it's a Google Drive link
        if "drive.google.com" in image_url:
            file_id = image_url.split("/d/")[1].split("/")[0]  # Extract File ID
            image_url = f"https://drive.google.com/uc?export=download&id={file_id}"
        
        headers = {
            "User-Agent": "Mozilla/5.0"
        }
        try:
            response = requests.get(image_url, headers=headers, stream=True)
        except Exception as e:
            response = requests.get("https://upload.wikimedia.org/wikipedia/commons/1/14/No_Image_Available.jpg", headers=headers, stream=True)
        
        if response.status_code!=200:
            response = requests.get("https://upload.wikimedia.org/wikipedia/commons/1/14/No_Image_Available.jpg", headers=headers, stream=True)
        
        if response.status_code == 200:
            os.makedirs(folder, exist_ok=True)  # Ensure folder exists
            file_path = os.path.join(folder, image_name)
            with open(file_path, "wb") as file:
                for chunk in response.iter_content(1024):
                    file.write(chunk)
            return image_name
        else:
            raise Exception(f"Failed to download image from {image_url}")

    except Exception as e:
        print(f"Error downloading image: {e}")

def insert_data_from_xl(excel_file: str)->bool:
    workbook = load_workbook(filename=excel_file, data_only=True)
    error=False
    error_message=[]
    c=1
    con1,cursor1 = get_con_cursor()

    try:
        q="""
            Select image from Products
        """
        cursor1.execute(q)
        images = cursor1.fetchall()
        for img in images:
            try:
                os.remove(f"app/media/{img[0]}")
            except Exception as e:
                print(e)
        q = """
            Delete From Products Where 1=1
        """
        cursor1.execute(q)
        con1.commit()
        q="""
            Delete From ProductCategory Where 1=1
        """
        cursor1.execute(q)
        con1.commit()
        q="""
            Delete From ProductSector Where 1=1
        """
        cursor1.execute(q)
        con1.commit()
    except Exception as e:
        con1.rollback()
        error=True
        error_message.append({
            "error":str(e)
        })
        return False
    finally:
        con1.close()
    
    category_cache = {}
    subcategory_cache = {}
    sector_cache = {}

    for sheet_name in workbook.sheetnames:
        sheet = workbook[sheet_name]

        for row in sheet.iter_rows(min_row=2):  # Skip headers
            con,cursor = get_con_cursor()
            try:
                name = row[0].value if row[0].value else None
                price = round(row[1].value,2) if row[1].value else 0.0
                types = row[2].value if row[2].value else "N/A"
                desc = row[3].value if row[3].value else "N/A"
                image_url = row[4].value if row[4].value else "https://upload.wikimedia.org/wikipedia/commons/1/14/No_Image_Available.jpg"
                status = row[5].value.lower() if row[5].value else "disable"
                category = row[6].value if row[6] else None
                subcategory = row[7].value if row[7] else None
                sector_name = row[8].value if row[8] else None
                qnt = row[9].value if row[9] else "N/A"

                if not category or not sector_name or not name:
                        continue
                
                # Handle sector name insertion if not exists
                if sector_name in sector_cache:
                    sector_id = sector_cache[sector_name]
                else:
                   cursor.execute("SELECT id FROM ProductSector WHERE sector_name = ?", (sector_name,)) 
                   sector = cursor.fetchone()
                   if sector:
                       sector_id = sector[0]
                   else:
                       sector_id = generate_unique_id()
                       cursor.execute("INSERT INTO ProductSector (id, sector_name) VALUES (?, ?)", (sector_id, sector_name))
                   sector_cache[sector_name] = sector_id

                # Handle category insertion if not exists
                if category in category_cache:
                    category_id = category_cache[category]
                else:
                    cursor.execute("SELECT id FROM ProductCategory WHERE category_name = ? AND NOT parent_category_id", (category,))
                    _category = cursor.fetchone()
                    if _category:
                        category_id = _category[0]
                    else:
                        category_id = generate_unique_id()
                        cursor.execute("INSERT INTO ProductCategory (id, category_name,sector_id) VALUES (?, ?, ?)", (category_id, category, sector_id))
                    category_cache[category] = category_id

                # Handle subcategory insertion if not exists
                if subcategory:
                    if (category, subcategory) in subcategory_cache:
                        subcategory_id = subcategory_cache[(category, subcategory)]
                    else:
                        cursor.execute(
                            "SELECT id FROM ProductCategory WHERE category_name = ? AND parent_category_id = ?",
                            (subcategory, category_id)
                        )
                        _subcategory = cursor.fetchone()
                        if _subcategory:
                            subcategory_id = _subcategory[0]
                        else:
                            subcategory_id = generate_unique_id()
                            cursor.execute(
                                "INSERT INTO ProductCategory (id, category_name, parent_category_id) VALUES (?, ?, ?)",
                                (subcategory_id, subcategory, category_id)
                            )
                        subcategory_cache[(category, subcategory)] = subcategory_id
                else:
                    subcategory_id = None

                p_id = generate_unique_id()
                image_name = p_id+".jpg"
                # img = download_save_image(
                #     image_url=image_url,
                #     image_name=image_name
                # )
                if "drive.google.com" in image_url:
                    file_id = image_url.split("/d/")[1].split("/")[0]  # Extract File ID
                    image_url = f"https://drive.google.com/thumbnail?id={file_id}"
                
                q = """
                    INSERT INTO Products (id, name, image, price, desc,type,is_active,price_gst_included,category_id,subcategory_id,qnt_unit)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """
                cursor.execute(q,(p_id,name,image_url,price,desc,types,1 if status=="enable" else 0,0,category_id,subcategory_id,qnt))
                con.commit()
                    
            except Exception as e:
                import traceback
                print(traceback.print_exc())
                print(e)
                con.rollback()
                error=True
                error_message.append({
                    "error":str(e),
                    "product":name,
                    "excel_row_no":c
                })
            finally:
                c+=1
                con.close()

        print("Sheet processing complete.")
        print("error : ",error_message)
        #con.close()
        category_cache={}
        subcategory_cache={}
        sector_cache={}
        if error:
            return False
        return True

def generate_unique_id():
    unique_id = uuid.uuid4()
    current_time_millis = int(time.time() * 1000)
    combined_id = str(current_time_millis) + str(unique_id).replace('-', '')
    return combined_id

def hashs(password:str)->str:
    try:
        password=password.encode()
        return bcrypt.hashpw(password, salt.encode("utf-8")).decode()
    except Exception as e:
        print(e)
        return ""

def verify_hash(password:str,hashed_password:str)->bool:
    try:
        password=password.encode()
        hashed_password=hashed_password.encode()
        return bcrypt.checkpw(password, hashed_password) 
    except Exception as e:
        print(e)
        return False

def send_email(subject: str, body: str, to: list = [], attachment=None):
    try:
        SMTP_SENDER = os.getenv('SMTP_SENDER')
        configuration = sib_api_v3_sdk.Configuration()
        configuration.api_key['api-key'] = os.getenv('BREVO_API_KEY')
        api_instance = sib_api_v3_sdk.TransactionalEmailsApi(sib_api_v3_sdk.ApiClient(configuration))

        send_smtp_email = SendSmtpEmail(
            to=[{"email": recipient} for recipient in to],
            sender={"email": SMTP_SENDER},
            subject=subject,
            html_content=body
        )

        if attachment:
            send_smtp_email.attachment = [
                {
                    "content": attachment['content'],
                    "name": attachment['name']
                }
            ]

        api_response = api_instance.send_transac_email(send_smtp_email)
        print(f"Email sent successfully: {api_response}")

    except ApiException as e:
        print(f"Exception when sending email: {e}")
    except Exception as e:
        print(f"An unexpected error occurred: {e}")

def send_email_smtp(subject: str, body: str, to: str, attachment=None,html=True):

    SMTP_USERNAME = os.getenv('SMTP_USERNAME')
    SMTP_SERVER = os.getenv('SMTP_SERVER')
    SMTP_PORT = os.getenv('SMTP_PORT')
    SMTP_PASSWORD = os.getenv('SMTP_PASSWORD')

    message = MIMEMultipart()
    message["From"] = SMTP_USERNAME
    message["To"] = to
    message["Subject"] = subject

    if html:
        message.attach(MIMEText(body, "html"))
    else:
        message.attach(MIMEText(body, "plain"))
    
    if attachment:
        try:
            attachment_content = base64.b64decode(attachment["content"])
            attachment_part = MIMEBase("application", "octet-stream")
            attachment_part.set_payload(attachment_content)
            encoders.encode_base64(attachment_part)
            attachment_part.add_header(
                "Content-Disposition",
                f"attachment; filename={attachment['name']}"
            )
            message.attach(attachment_part)
        except Exception as e:
            print(f"Error attaching file: {e}")

    try:
        # ✅ Use SMTP_SSL for port 465
        with smtplib.SMTP_SSL(SMTP_SERVER, SMTP_PORT) as server:
            server.login(SMTP_USERNAME, SMTP_PASSWORD)
            server.sendmail(SMTP_USERNAME, to, message.as_string())
            print("Email sent successfully")
    except Exception as e:
        print(f"Email send error: {e}")


def push_email(template_addr: str, receiver_email: list, subject: str, context: dict = {},attachment=None):
    try:
        template = transform(render_template(template_addr, **context) if context else render_template(template_addr))

        p = Process(
            target=send_email_smtp,
            args=(subject, template, receiver_email[0],attachment)
        )
        p.start()
    except Exception as e:
        print(f"push_email() | {str(e)}")


if __name__=="__main__":
    # pswd=hashs("utsav")
    # print(pswd)
    # h_p="$2b$12$Dd8MfO0BvaaajOPoh3A7SOVG8S3zbIp3PWHfjHLWVidnxAf4ELrFS"
    # pswd="utsav"
    # print(verify_hash(pswd,h_p))
    # read_xl('test.xlsx','test')
    # insert_data_from_xl("app/media/importedExcels/1738234101145ssp product structure (2).xlsx")
    # download_save_image('test.jpg','https://drive.google.com/file/d/1ISH4ZOt7WM_M9vvpTNJYK8D7cDcVMplM/view?usp=sharing')
    pass